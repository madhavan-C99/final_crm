import os
import tempfile
from unittest.mock import patch
from types import SimpleNamespace

import openpyxl
from django.test import SimpleTestCase, override_settings
from rest_framework.exceptions import PermissionDenied

from adm.services.query_services import replace_query
from adm.views.query_views import _scope_query_filters
from adm.services.generic_export_services import (
    _fetch_export_rows,
    export_data_service,
)


class GenericExportServiceTests(SimpleTestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.settings_override = override_settings(MEDIA_ROOT=self.temp_dir.name)
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.user = object()

    @patch("adm.services.lead_services.fetch_all_leads_admin")
    def test_leads_export_uses_all_filtered_leads(self, fetch_leads):
        fetch_leads.return_value = {"leads": [{"id": 7}]}

        rows = _fetch_export_rows("leads", self.user, {"search": "Asha"})

        self.assertEqual(rows, [{"id": 7}])
        fetch_leads.assert_called_once_with(
            user=self.user,
            search="Asha",
            page_size="all",
        )

    @patch("adm.services.payment_services.fetch_all_pending_payments_admin")
    def test_pending_payment_export_loads_all_filtered_rows(self, fetch_payments):
        fetch_payments.return_value = {"leads": [{"id": 8}]}

        rows = _fetch_export_rows("pending_payments", self.user, {"search": "Maya"})

        self.assertEqual(rows, [{"id": 8}])
        fetch_payments.assert_called_once_with(
            user=self.user,
            all_rows=True,
            search="Maya",
            page=1,
        )

    @patch(
        "adm.services.loss_lead_approval_services."
        "fetch_loss_lead_approval_requests_admin"
    )
    def test_loss_approval_export_loads_all_filtered_rows(self, fetch_approvals):
        fetch_approvals.return_value = {"data": {"leads": [{"lead_id": 9}]}}

        rows = _fetch_export_rows("loss_approvals", self.user, {"search": "Ravi"})

        self.assertEqual(rows, [{"lead_id": 9}])
        fetch_approvals.assert_called_once_with(
            user=self.user,
            search="Ravi",
            page_size="all",
        )

    @patch("adm.services.performance_services.fetch_performance_overview_admin")
    def test_performance_export_loads_all_filtered_rows(self, fetch_performance):
        fetch_performance.return_value = {
            "data": {"performance_list": [{"telecaller_id": 10}]}
        }

        rows = _fetch_export_rows("performance", self.user, {"team_id": 3})

        self.assertEqual(rows, [{"telecaller_id": 10}])
        fetch_performance.assert_called_once_with(
            {"team_id": 3, "page": 1, "page_size": 1},
            user=self.user,
            all_rows=True,
        )

    @patch(
        "adm.services.generic_export_services._fetch_export_rows",
        return_value=[
            {"id": 7, "full_name": "Asha Kumar", "mobile_no": "9876543210"},
            {"id": 8, "full_name": "Maya Raj", "mobile_no": "9123456780"},
        ],
    )
    def test_excel_export_has_green_header_spacing_and_selected_ids(
        self, fetch_rows
    ):
        result = export_data_service(
            user=self.user,
            entity="leads",
            selected_ids=[8],
            columns=["id", "full_name", "mobile_no"],
        )

        workbook_path = os.path.join(
            self.temp_dir.name,
            "exports",
            result["file_name"],
        )
        workbook = openpyxl.load_workbook(workbook_path)
        worksheet = workbook.active

        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            ["Id", "Full Name", "Mobile No"],
        )
        self.assertTrue(worksheet["A1"].fill.fgColor.rgb.endswith("84C225"))
        self.assertGreaterEqual(worksheet.column_dimensions["A"].width, 18)
        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual(worksheet["A2"].value, 8)
        self.assertEqual(result["total_exported"], 1)
        fetch_rows.assert_called_once_with("leads", self.user, {})

    @patch(
        "adm.services.generic_export_services._fetch_export_rows",
        return_value=[{"telecaller_id": 10, "telecaller_name": "Ravi"}],
    )
    def test_selected_ids_use_entity_identifier(self, fetch_rows):
        result = export_data_service(
            user=self.user,
            entity="performance",
            selected_ids=[10],
            columns=["telecaller_id", "telecaller_name"],
        )

        workbook_path = os.path.join(
            self.temp_dir.name,
            "exports",
            result["file_name"],
        )
        worksheet = openpyxl.load_workbook(workbook_path).active
        self.assertEqual(worksheet["A2"].value, 10)
        self.assertEqual(result["total_exported"], 1)


class OrganizationQueryScopeTests(SimpleTestCase):
    def test_raw_query_scope_does_not_include_unassigned_or_global_rows(self):
        query = (
            "SELECT * FROM telecalling_lead "
            "WHERE current_status = 'working' OR current_status = 'new' "
            "ORDER BY id"
        )

        scoped_query = replace_query(query, {"organization_id": 17})

        self.assertIn(
            "(current_status = 'working' OR current_status = 'new') "
            "AND telecalling_lead.organization_id = 17 ORDER BY id",
            scoped_query,
        )
        self.assertNotIn("organization_id IS NULL", scoped_query)
        self.assertNotIn("organization_id = 0", scoped_query)

    def test_existing_organization_placeholder_cannot_broaden_scope(self):
        query = (
            "SELECT * FROM telecalling_lead "
            "WHERE organization_id = @_organization_id OR @_organization_id = 0"
        )

        scoped_query = replace_query(query, {"organization_id": 17})

        self.assertIn("telecalling_lead.organization_id = 17", scoped_query)

    def test_zero_organization_id_does_not_disable_tenant_filter(self):
        query = "SELECT * FROM telecalling_lead"

        scoped_query = replace_query(query, {"organization_id": 0})

        self.assertIn("telecalling_lead.organization_id = 0", scoped_query)

    def test_regular_user_cannot_override_their_organization_filter(self):
        user = SimpleNamespace(id=5, organization_id=17, is_superuser=False)

        filters = _scope_query_filters(
            user,
            {"organization_id": 99, "user_id": 99, "status": "open"},
        )

        self.assertEqual(filters, {"organization_id": 17, "user_id": 5, "status": "open"})

    def test_user_without_organization_is_denied(self):
        user = SimpleNamespace(id=5, organization_id=None, is_superuser=False)

        with self.assertRaises(PermissionDenied):
            _scope_query_filters(user, {})

    def test_platform_super_admin_can_use_unscoped_filter(self):
        user = SimpleNamespace(id=1, organization_id=None, is_superuser=True)

        self.assertEqual(_scope_query_filters(user, {"organization_id": 0}), {})
