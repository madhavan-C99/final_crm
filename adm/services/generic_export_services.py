import os
import csv
import uuid
from datetime import datetime
from django.conf import settings
from rest_framework.exceptions import APIException
from telecalling.models import User
from adm.models import Team
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date


from .lead_services import fetch_all_leads_admin
from .payment_services import fetch_all_pending_payments_admin
from .loss_lead_approval_services import fetch_loss_lead_approval_requests_admin
from .performance_services import fetch_performance_overview_admin

FULL_ADMIN_ROLE_CODES = {"ADM", "MGR", "AMGR", "DEV", "SUPERADMIN", "ADMIN"}
TEAM_LEAD_ROLE_CODES = {"TEAM_LEAD", "TL"}

HEADER_DISPLAY_NAMES = {
    "s_no": "S.No",
    "full_name": "Name",
    "mobile_no": "Mobile No",
    "pipeline_name": "Pipeline",
    "pipeline_stage_name": "Stage",
    "assigned_to": "Assigned Telecallers",
    "priority_name": "Priority",
    "campaign_name": "Campaign",
    "source_name": "Source",
    "course_name": "Course Name",
    "course_plan_name": "Course Plan",
    "amount": "Amount",
    "pending_amount": "Pending Amount",
    "next_followup": "Next Followup",
    "last_contacted": "Last Contacted",
    "enquiry_date": "Enquiry Date",
    "user_name": "User Name",
    "reporting_manager": "Reporting Manager",
    "mobile_number": "Mobile Number",
    "date": "Date",
    "total_disposed_count": "Total Disposed Count",
    "disposed_yes_connected_count": "Disposed Yes Connected Count",
    "disposed_not_connected_count": "Disposed Not Connected Count",
    "total_followup_leads": "Total Followup Leads",
    "total_converted_leads": "Total Converted Leads",
    "total_lost_leads": "Total Lost Leads",
    "conversion_percentage": "Conversion %",
    "total_assigned_leads": "Total Assigned Leads",
    "followups_due_today": "Followups Due Today",
    "average_tat": "Average TAT",
    "campaign_name": "Campaign Name",
    "pipeline_name": "Pipeline Name",
    "total_lead_count": "Total Lead Count",
    "total_leads": "Total Leads",
    "total_unassigned_leads": "Total Unassigned Leads",
    "total_open_leads": "Total Open Leads",
    "total_in_progress_leads": "Total In-Progress Lead",
    "total_closed_leads": "Total Closed Leads",
    "total_disconnected_lost_leads": "Total Disconnected Lead (Lost)",
    "total_disconnected_pending_leads": "Total Disconnected (Pending)",
    "total_connected_converted_leads": "Total Connected Leads (Converted/Joined)",
    "total_connected_lost_leads": "Total Connected Leads (Lost/Closed)",
    "total_calls_connected": "Total Calls Connected",
    "leads_marked_as_converted": "Leads Marked As Converted",
    "leads_marked_as_lost": "Leads Marked As Lost",
    "date_range": "Date Range",
    "total_joined_converted": "Total Joined / Converted",
    "won_value": "Won Value",
    "total_lost_value": "Total Lost Value",
}


def get_user_role_code(user):
    user_role_obj = user.user_roles.select_related('role').first() if hasattr(user, 'user_roles') else None
    if user_role_obj and user_role_obj.role and user_role_obj.role.code:
        return user_role_obj.role.code

    if getattr(user, 'is_superuser', False):
        return "ADM"

    return ""


def is_admin_or_manager_user(user):
    role_code = get_user_role_code(user)
    return role_code in FULL_ADMIN_ROLE_CODES or role_code in TEAM_LEAD_ROLE_CODES


def get_team_user_ids(user):
    user_ids = {user.id}
    if hasattr(user, 'team') and user.team:
        member_ids = set(User.objects.filter(team=user.team).values_list('id', flat=True))
        user_ids.update(member_ids)
    
    led_team_ids = set(Team.objects.filter(leader=user).values_list('id', flat=True))
    if led_team_ids:
        led_member_ids = set(User.objects.filter(team_id__in=led_team_ids).values_list('id', flat=True))
        user_ids.update(led_member_ids)

    return list(user_ids)


def apply_role_data_scoping(user, entity, filters):
    role_code = get_user_role_code(user)

    if role_code in TEAM_LEAD_ROLE_CODES:
        key = "assigned_to_ids" if entity in ["leads", "pending_payments"] else "telecaller_ids"
        filters[key] = get_team_user_ids(user)
    elif role_code not in FULL_ADMIN_ROLE_CODES:
        key = "assigned_to" if entity in ["leads", "pending_payments"] else "telecaller_id"
        filters[key] = user.id

    return filters


def fetch_export_records(entity, user, filters):
    filters = apply_role_data_scoping(user, entity, filters)

    if entity == "leads":
        filters["page_size"] = "all"
        return fetch_all_leads_admin(user=user, **filters).get("leads", [])

    if entity == "pending_payments":
        filters.pop("all_rows", None)
        filters["page"] = 1
        return fetch_all_pending_payments_admin(user=user, all_rows=True, **filters).get("leads", [])

    if entity == "loss_approvals":
        filters["page_size"] = "all"
        res = fetch_loss_lead_approval_requests_admin(user=user, **filters)
        if isinstance(res, dict):
            return res.get("requests") or res.get("leads") or res.get("data", {}).get("leads", [])
        return []

    if entity == "performance":
        filters["page"] = 1
        filters["page_size"] = 1
        return fetch_performance_overview_admin(filters, user=user, all_rows=True).get("data", {}).get("performance_list", [])

    if entity == "reports":
        from .report_services import execute_report_service
        report_payload = dict(filters or {})
        report_payload["page"] = 1
        report_payload["limit"] = 10000
        res = execute_report_service(report_payload)
        return res.get("data", [])

    raise APIException(f"Unsupported export entity: {entity}")


COLUMN_SYNONYMS = {
    ('name', 'full_name', 'contact_name', 'student_name'): ['full_name', 'name', 'telecaller_name'],
    ('mobile', 'mobile_no', 'contact', 'phone', 'number'): ['mobile_no', 'contact', 'mobile'],
    ('stage', 'pipeline_stage', 'stage_name', 'pipeline_stage_name', 'status'): ['pipeline_stage_name', 'stage', 'approval_status', 'status'],
    ('pipeline', 'pipeline_name'): ['pipeline_name', 'pipeline'],
    ('priority', 'priority_name', 'tag'): ['priority_name', 'tag'],
    ('campaign', 'campaign_name'): ['campaign_name', 'campaign'],
    ('source', 'source_name', 'lead_source'): ['source_name', 'lead_source', 'source'],
    ('course', 'course_name'): ['course_name', 'course'],
    ('course_plan', 'course_plan_name', 'plan'): ['course_plan_name', 'course_plan'],
    ('next_follow_up', 'next_followup'): ['next_followup'],
    ('created', 'enquiry', 'joining', 'inquiry', 'date', 'enquiry_date'): ['enquiry_date', 'joining_date', 'inquiry_date', 'created_at'],
}


def resolve_export_columns(requested_cols, available_keys):
    if not requested_cols:
        return available_keys

    matched_cols = []
    for req in requested_cols:
        req_clean = str(req).strip().lower().replace(" ", "_")
        if req_clean in available_keys:
            if req_clean not in matched_cols:
                matched_cols.append(req_clean)
            continue

        found = False
        for keys_tuple, target_list in COLUMN_SYNONYMS.items():
            if req_clean in keys_tuple:
                for target in target_list:
                    if target in available_keys and target not in matched_cols:
                        matched_cols.append(target)
                        found = True
                        break
            if found:
                break

        if not found:
            for avail in available_keys:
                if req_clean in avail.lower() and avail not in matched_cols:
                    matched_cols.append(avail)
                    break

    return matched_cols if matched_cols else available_keys



def generate_excel_file(file_path, export_cols, raw_rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Export Data"

    headers = [HEADER_DISPLAY_NAMES.get(col, str(col).replace("_", " ").title()) for col in export_cols]
    ws.append(headers)

    for row in raw_rows:
        clean_row = []
        for col in export_cols:
            val = row.get(col, "")
            if hasattr(val, "tzinfo") and getattr(val, "tzinfo", None) is not None:
                val = val.replace(tzinfo=None)
            elif isinstance(val, (datetime, date)):
                val = val.strftime("%d %b %Y, %I:%M %p") if isinstance(val, datetime) else val.strftime("%d %b %Y")
            elif val is None:
                val = ""
            clean_row.append(val)
        ws.append(clean_row)

    header_fill = PatternFill(fill_type="solid", fgColor="84C225")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # Style Header Row
    ws.row_dimensions[1].height = 28
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    # Style Data Rows
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 24
        for cell in row:
            cell.font = data_font
            cell.alignment = center_align
            cell.border = border

    # 🟢 Auto-Calculate Generous Column Widths (Prevents '###' and cramped text)
    for col_cells in ws.columns:
        max_length = 0
        for cell in col_cells:
            val_str = str(cell.value or "")
            if val_str:
                max_length = max(max_length, len(val_str))
        
        column_letter = get_column_letter(col_cells[0].column)
        ws.column_dimensions[column_letter].width = min(max(max_length + 6, 20), 60)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(file_path)


def export_data_service(user, entity, export_format, selected_columns, filters=None):
    try:
        filters = filters or {}

        raw_rows = fetch_export_records(entity, user, filters)
        print(f"Raw rows for export: {len(raw_rows)}")
        if not raw_rows:
            return {
                "status": "success",
                "message": "No records found to export",
                "total_exported": 0,
                "download_url": ""
            }

        available_keys = list(raw_rows[0].keys())
        export_cols = resolve_export_columns(selected_columns, available_keys)

        media_export_dir = os.path.join(settings.MEDIA_ROOT, 'exports')
        os.makedirs(media_export_dir, exist_ok=True)
        report_key = str(filters.get("report_key") or "").strip()
        base_name = report_key if (entity == "reports" and report_key) else entity
        filename = f"{base_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}.xlsx"
        file_path = os.path.join(media_export_dir, filename)

        generate_excel_file(file_path, export_cols, raw_rows)

        if entity == "reports" and report_key:
            formatted_name = "User Follow Up Report" if "followup" in report_key.lower() else base_name.replace('_', ' ').title()
            custom_message = f"{formatted_name} exported successfully"
        else:
            custom_message = f"Successfully exported {len(raw_rows)} records in {export_format.upper()} format!"

        return {
            "status": "success",
            "message": custom_message,
            "export_format": export_format,
            "total_exported": len(raw_rows),
            "file_name": filename,
            "download_url": f"{settings.MEDIA_URL}exports/{filename}"
        }

    except Exception as e:
        if isinstance(e, APIException):
            raise
        raise APIException(str(e)) from e
