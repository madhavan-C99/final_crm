from .query_services import exec_raw_sql
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.db.models import (
    DateField,
    DateTimeField,
    F,
    OuterRef,
    Q,
    Subquery,
    Sum,
    Count,
)
from django.db.models.functions import Coalesce, TruncDate
from django.utils import timezone
from datetime import datetime, date
from rest_framework.exceptions import APIException

from telecalling.models import (
    Lead, PaymentInfo, PaymentHistory, FollowUp, CallDetails,
    PipelineStage, User, CourseName, CoursePlan, CourseTiming, PaymentStage
)

from datetime import datetime, date, timedelta

def get_pending_payment_filter_dropdowns_admin(user=None):
    """
    Pending Payments Page -> Filter Modal Dropdowns API.
    Fast Dropdowns using Collection Queries (exec_raw_sql).
    """
    try:
        org_id = getattr(user, 'organization_id', 0) if user else 0
        params = {'organization_id': org_id}

        course_names = exec_raw_sql('L_COURSE_NAMES', params)
        course_plans = exec_raw_sql('L_COURSE_PLANS', params)
        course_timings = exec_raw_sql('L_COURSE_TIMINGS', params)

        payment_stages = [
            {"value": "today_due", "label": "Today Due"},
            {"value": "active_due", "label": "Active Due"},
            {"value": "overdue", "label": "Overdue"}
        ]

        return {
            "status": "success",
            "data": {
                "course_names": course_names,
                "course_plans": course_plans,
                "course_timings": course_timings,
                "payment_stages": payment_stages
            }
        }
    except Exception as e:
        raise APIException(str(e))


def fetch_all_pending_payments_admin(
    user=None,
    search=None,
    date_filter=None,
    from_date=None,
    to_date=None,
    date_filter_type=None,
    sort_by=None,
    sort_type=None,
    pipeline_id=None,
    course_name_id=None,
    course_plan_id=None,
    course_timing_id=None,
    payment_stage_id=None,
    pending_amount_range=None,
    page=1,
    page_size=None,
    limit=None,
    all_rows=False,
    **kwargs,
):
    try:
        today = timezone.localdate()
        page = max(1, int(page or 1))
        page_size = int(page_size or limit or 50)
        if page_size < 1:
            raise ValueError("page_size must be greater than zero")

        course_name = kwargs.get("course_name") or course_name_id
        course_plan = kwargs.get("course_plan") or course_plan_id
        course_time = kwargs.get("course_time") or course_timing_id
        payment_stage = kwargs.get("payment_stage") or payment_stage_id
        pending_amount = kwargs.get("pending_amount") or pending_amount_range

        latest_payment_due = PaymentHistory.objects.filter(
            payment_id=OuterRef("pk")
        ).order_by("-id").values("due_date")[:1]
        next_followup = FollowUp.objects.filter(
            lead_id=OuterRef("lead_id"),
            is_attended=False,
            scheduled_at__isnull=False,
        ).order_by("scheduled_at")
        latest_call = CallDetails.objects.filter(
            lead_id=OuterRef("lead_id")
        ).order_by("-called_at")

        payments_qs = PaymentInfo.objects.select_related(
            "lead",
            "lead__assigned_to",
            "lead__pipeline_stage",
            "lead__campaign",
            "lead__course_plan",
            "lead__course_name",
            "lead__course",
            "lead__course__name",
            "lead__course_timing",
        ).filter(
            pending_amount__gt=0,
        ).annotate(
            resolved_due_date=Coalesce(
                Subquery(latest_payment_due, output_field=DateField()),
                TruncDate(
                    Subquery(
                        next_followup.values("scheduled_at")[:1],
                        output_field=DateTimeField(),
                    )
                ),
                TruncDate(F("lead__created_at")),
                output_field=DateField(),
            ),
            next_followup_at=Subquery(
                next_followup.values("scheduled_at")[:1],
                output_field=DateTimeField(),
            ),
            next_followup_notes=Subquery(
                next_followup.values("notes")[:1],
            ),
            last_call_at=Subquery(
                latest_call.values("called_at")[:1],
                output_field=DateTimeField(),
            ),
            last_call_summary=Subquery(
                latest_call.values("conversation_summary")[:1],
            ),
        )

        if user and getattr(user, "is_authenticated", False) and getattr(user, "organization", None):
            payments_qs = payments_qs.filter(lead__organization=user.organization)
        elif user and getattr(user, "is_authenticated", False):
            payments_qs = payments_qs.none()

        if search and str(search).strip():
            search = str(search).strip()
            payments_qs = payments_qs.filter(
                Q(lead__full_name__icontains=search)
                | Q(lead__mobile_no__icontains=search)
                | Q(lead__email__icontains=search)
            )

        if pipeline_id:
            payments_qs = payments_qs.filter(
                Q(lead__pipeline_stage__pipeline_category_id=pipeline_id)
                | Q(lead__campaign__pipeline_category_id=pipeline_id)
            )

        if course_name:
            course_name_value = str(course_name).strip()
            if course_name_value.isdigit():
                payments_qs = payments_qs.filter(
                    Q(lead__course_name_id=int(course_name_value))
                    | Q(lead__course__name_id=int(course_name_value))
                )
            else:
                payments_qs = payments_qs.filter(
                    Q(lead__course_name__coursename__icontains=course_name_value)
                    | Q(lead__course__name__coursename__icontains=course_name_value)
                )

        if course_plan:
            course_plan_value = str(course_plan).strip()
            if course_plan_value.isdigit():
                payments_qs = payments_qs.filter(
                    Q(lead__course_plan_id=int(course_plan_value))
                    | Q(lead__course__plan_id=int(course_plan_value))
                )
            else:
                payments_qs = payments_qs.filter(
                    Q(lead__course_plan__courseplan__icontains=course_plan_value)
                    | Q(lead__course__plan__courseplan__icontains=course_plan_value)
                )

        if course_time:
            course_time_value = str(course_time).strip()
            if course_time_value.isdigit():
                payments_qs = payments_qs.filter(
                    Q(lead__course_timing_id=int(course_time_value))
                    | Q(lead__course__time_id=int(course_time_value))
                )
            else:
                payments_qs = payments_qs.filter(
                    Q(lead__course_timing__coursetime__icontains=course_time_value)
                    | Q(lead__course__time__coursetime__icontains=course_time_value)
                )

        amount_filter = str(pending_amount or "").lower().strip()
        if amount_filter in {"above_5k", "1", "above_5000"}:
            payments_qs = payments_qs.filter(pending_amount__gt=5000)
        elif amount_filter in {"below_5k", "2", "below_5000"}:
            payments_qs = payments_qs.filter(pending_amount__lte=5000)
        elif amount_filter in {"below_2k", "3", "below_2000"}:
            payments_qs = payments_qs.filter(pending_amount__lte=2000)
        elif amount_filter == "5000_10000":
            payments_qs = payments_qs.filter(
                pending_amount__gte=5000, pending_amount__lte=10000
            )
        elif amount_filter == "above_10000":
            payments_qs = payments_qs.filter(pending_amount__gt=10000)

        summary = payments_qs.aggregate(
            total_amount=Sum("pending_amount"),
            total_count=Count("id"),
            due_today_amount=Sum(
                "pending_amount",
                filter=Q(resolved_due_date=today),
            ),
            due_today_count=Count(
                "id",
                filter=Q(resolved_due_date=today),
            ),
            overdue_amount=Sum(
                "pending_amount",
                filter=Q(resolved_due_date__lt=today),
            ),
            overdue_count=Count(
                "id",
                filter=Q(resolved_due_date__lt=today),
            ),
        )
        total_pending_amount = round(float(summary["total_amount"] or 0), 2)
        total_pending_leads = int(summary["total_count"] or 0)
        due_today_amount = round(float(summary["due_today_amount"] or 0), 2)
        due_today_leads = int(summary["due_today_count"] or 0)
        overdue_amount = round(float(summary["overdue_amount"] or 0), 2)
        overdue_leads = int(summary["overdue_count"] or 0)

        filtered_qs = payments_qs
        filter_type = str(date_filter_type or date_filter or "").lower().strip()
        filter_from = None
        filter_to = None
        if filter_type == "today":
            filter_from = filter_to = today
        elif filter_type == "yesterday":
            filter_from = filter_to = today - timedelta(days=1)
        elif filter_type in {"this_week", "weekly", "week"}:
            filter_from = today - timedelta(days=today.weekday())
            filter_to = filter_from + timedelta(days=6)
        elif filter_type in {"this_month", "monthly", "month"}:
            filter_from = today.replace(day=1)
            filter_to = (today.replace(day=28) + timedelta(days=4))
            filter_to = filter_to - timedelta(days=filter_to.day)
        elif filter_type in {"this_year", "yearly", "year"}:
            filter_from = date(today.year, 1, 1)
            filter_to = date(today.year, 12, 31)
        elif filter_type == "custom" or from_date or to_date:
            filter_from = from_date
            filter_to = to_date

        if filter_type == "overdue":
            filtered_qs = filtered_qs.filter(resolved_due_date__lt=today)
        elif filter_type == "upcoming":
            filtered_qs = filtered_qs.filter(resolved_due_date__gt=today)
        elif filter_from or filter_to:
            if filter_from:
                filtered_qs = filtered_qs.filter(resolved_due_date__gte=filter_from)
            if filter_to:
                filtered_qs = filtered_qs.filter(resolved_due_date__lte=filter_to)

        payment_stage_value = str(payment_stage or "").lower().strip().replace(" ", "_")
        if payment_stage_value in {"today_due", "due_today"}:
            filtered_qs = filtered_qs.filter(resolved_due_date=today)
        elif payment_stage_value in {"active_due", "active"}:
            filtered_qs = filtered_qs.filter(
                Q(resolved_due_date__gt=today) | Q(resolved_due_date__isnull=True)
            )
        elif payment_stage_value == "overdue":
            filtered_qs = filtered_qs.filter(resolved_due_date__lt=today)
        elif payment_stage_value.isdigit():
            latest_payment_stage = PaymentHistory.objects.filter(
                payment_id=OuterRef("pk")
            ).order_by("-id").values("due_stage_id")[:1]
            filtered_qs = filtered_qs.annotate(
                latest_payment_stage_id=Subquery(latest_payment_stage)
            ).filter(latest_payment_stage_id=int(payment_stage_value))
        elif payment_stage_value:
            filtered_qs = filtered_qs.filter(
                payment_histories__due_stage__name__iexact=payment_stage
            ).distinct()

        total_records = filtered_qs.count()
        total_pages = (total_records + page_size - 1) // page_size
        if sort_by in {"due_date_asc", "due_date_desc"}:
            ordering = "resolved_due_date" if sort_by.endswith("_asc") else "-resolved_due_date"
        elif sort_by in {"amount_asc", "amount_desc"}:
            ordering = "pending_amount" if sort_by.endswith("_asc") else "-pending_amount"
        else:
            order_type = str(sort_type or kwargs.get("sort_type") or "").lower().strip()
            if order_type not in {"newest", "oldest"}:
                order_type = "oldest" if sort_by in {"oldest", "created_at"} else "newest"
            ordering = "id" if order_type == "oldest" else "-id"

        if all_rows:
            start = 0
            page_qs = filtered_qs.order_by(ordering)
        else:
            start = (page - 1) * page_size
            page_qs = filtered_qs.order_by(ordering)[start:start + page_size]

        processed_leads = []
        for p in page_qs:
            lead = p.lead
            pending_amt = float(p.pending_amount or 0)
            due_date = p.resolved_due_date
            if due_date and due_date < today:
                status_text = "Overdue"
            elif due_date == today:
                status_text = "Due Today"
            else:
                status_text = "Active"

            assigned_name = "-"
            if lead.assigned_to:
                first_name = (lead.assigned_to.first_name or "").strip()
                last_name = (lead.assigned_to.last_name or "").strip()
                assigned_name = (
                    f"{first_name} {last_name}".strip()
                    or lead.assigned_to.username
                    or "-"
                )

            last_conversation = p.last_call_summary or p.next_followup_notes
            if not last_conversation:
                last_conversation = "No conversation recorded yet"

            processed_leads.append({
                "s_no": start + len(processed_leads) + 1,
                "id": lead.id,
                "payment_id": p.id,
                "name": lead.full_name or "-",
                "contact": lead.mobile_no or "-",
                "email": lead.email or "-",
                "assigned_to": assigned_name,
                "pipeline": lead.pipeline_stage.name.title() if lead.pipeline_stage else "-",
                "campaign": lead.campaign.name if lead.campaign else "-",
                "course_plan": getattr(lead.course_plan, "courseplan", "-") if lead.course_plan else "-",
                "course": (
                    getattr(lead.course_name, "coursename", None)
                    or (
                        getattr(lead.course.name, "coursename", "-")
                        if lead.course and lead.course.name else "-"
                    )
                ),
                "joining_date": (
                    lead.created_at.strftime("%d %b, %I:%M %p")
                    if lead.created_at else "-"
                ),
                "batch_timing": (
                    getattr(lead.course_timing, "coursetime", "-")
                    if lead.course_timing else "Standard Batch"
                ),
                "amount_paid": float(p.amount_paid or 0),
                "pending_amount": pending_amt,
                "status": status_text,
                "due_date": due_date.isoformat() if due_date else None,
                "next_followup": (
                    timezone.localtime(p.next_followup_at).strftime("%d %b, %I:%M %p")
                    if p.next_followup_at else None
                ),
                "last_conversation": last_conversation,
            })

        summary_cards = {
            "total_pending": {
                "amount": total_pending_amount,
                "count": total_pending_leads,
            },
            "due_today": {
                "amount": due_today_amount,
                "count": due_today_leads,
            },
            "overdue": {
                "amount": overdue_amount,
                "count": overdue_leads,
            },
        }

        return {
            "status": True,
            "total_records": total_records,
            "total_pages": total_pages,
            "summary_cards": summary_cards,
            "leads": processed_leads,
            "total_pending_amount": total_pending_amount,
            "total_leads": total_pending_leads,
            "today_due_amount": due_today_amount,
            "today_due_leads": due_today_leads,
            "due_today_amount": due_today_amount,
            "due_today_leads": due_today_leads,
            "overdue_amount": overdue_amount,
            "overdue_leads": overdue_leads,
            "total_count": total_records,
            "page": page,
            "page_size": page_size,
            "limit": page_size,
        }

    except Exception as e:
        raise APIException(str(e))


def export_pending_payments_admin(
    search=None, date_filter=None, from_date=None, to_date=None, date_filter_type=None,
    sort_by=None, pipeline_id=None, course_name_id=None, course_plan_id=None,
    course_timing_id=None, payment_stage_id=None, pending_amount_range=None, **kwargs
):
 
    try:
        today = timezone.now().date()

        # Fetch pending payments data using existing service logic
        res = fetch_all_pending_payments_admin(
            search=search,
            date_filter=date_filter,
            from_date=from_date,
            to_date=to_date,
            date_filter_type=date_filter_type,
            sort_by=sort_by,
            pipeline_id=pipeline_id,
            course_name_id=course_name_id,
            course_plan_id=course_plan_id,
            course_timing_id=course_timing_id,
            payment_stage_id=payment_stage_id,
            page=1,
            limit=10000
        )
        leads = res.get("leads", [])

        # Create OpenPyXL Workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Pending Payments"
        ws.views.sheetView[0].showGridLines = True

        # Styles definition
        header_fill = PatternFill(start_color="84C225", end_color="84C225", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        data_font = Font(name="Calibri", size=10, color="000000")
        data_align = Alignment(horizontal="center", vertical="center")

        thin_side = Side(border_style="thin", color="D3D3D3")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # 17 Columns Headers
        headers = [
            "S.No",
            "Student Name",
            "Contact Number",
            "Email Address",
            "Assigned Telecaller",
            "Pipeline Stage",
            "Campaign",
            "Course Plan",
            "Course Name",
            "Joining Date",
            "Batch & Timing",
            "Amount Paid (₹)",
            "Pending Amount (₹)",
            "Payment Status",
            "Due Date",
            "Next Followup",
            "Last Conversation Summary"
        ]

        ws.row_dimensions[1].height = 32
        ws.append(headers)

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border

        # Add Data Rows
        for row_idx, lead in enumerate(leads, start=2):
            row_data = [
                lead.get("s_no", row_idx - 1),
                lead.get("name", "-"),
                lead.get("contact", "-"),
                lead.get("email") or "-",
                lead.get("assigned_to", "-"),
                lead.get("pipeline", "-"),
                lead.get("campaign", "-"),
                lead.get("course_plan", "-"),
                lead.get("course", "-"),
                lead.get("joining_date", "-"),
                lead.get("batch_timing", "-"),
                lead.get("amount_paid", 0.0),
                lead.get("pending_amount", 0.0),
                lead.get("status", "-"),
                lead.get("due_date") or "-",
                lead.get("next_followup") or "-",
                lead.get("last_conversation", "-")
            ]

            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 25

            for col_idx in range(1, len(row_data) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.alignment = data_align
                cell.border = thin_border

        # Auto-adjust Column Widths (+6 padding)
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 6, 15)

        # Create HTTP Response for Excel Download
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="Pending_Payments_Report_{today.strftime("%Y%m%d")}.xlsx"'
        wb.save(response)
        return response

    except Exception as e:
        raise APIException(str(e))
